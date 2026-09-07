"""Find WHERE inside a deck or PDF the query best matches - the slide or page number.

Simple and fast: score each slide/page by how many query words it contains,
and return the best one. (FR-5: 'the precise spot - slide/page of the match'.)
"""
from pathlib import Path


def _score(text, qwords):
    t = (text or "").lower()
    return sum(1 for w in qwords if w in t)


def _snip(text, n=160):
    text = " ".join((text or "").split())
    return (text[:n] + "...") if len(text) > n else text


def locate_in_transcript(segments, query):
    """segments: list of [start_seconds, text]. Return (seconds:int, snippet) of the
    best-matching moment, or (None, '') if nothing matches. Powers the video jump."""
    qwords = [w for w in query.lower().split() if len(w) > 1]
    if not qwords:
        return None, ""
    best = (0, None, "")  # score, start, text
    for start, text in segments:
        s = _score(text, qwords)
        if s > best[0]:
            best = (s, start, text)
    if best[1] is not None and best[0] > 0:
        return int(best[1]), _snip(best[2])
    return None, ""


def locate_in_file(abs_path, query):
    """Return (where_label, snippet), e.g. ('Slide 4', '...'). Empty strings if unknown."""
    p = Path(abs_path)
    ext = p.suffix.lower()
    qwords = [w for w in query.lower().split() if len(w) > 1]
    if not qwords:
        return "", ""
    try:
        if ext in (".pptx", ".ppt"):
            from pptx import Presentation
            prs = Presentation(str(p))
            best = (0, 0, "")
            for i, slide in enumerate(prs.slides, start=1):
                txt = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
                s = _score(txt, qwords)
                if s > best[0]:
                    best = (s, i, txt)
            if best[0] > 0:
                return f"Slide {best[1]}", _snip(best[2])
        elif ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(p))
            best = (0, 0, "")
            for i, page in enumerate(reader.pages, start=1):
                txt = page.extract_text() or ""
                s = _score(txt, qwords)
                if s > best[0]:
                    best = (s, i, txt)
            if best[0] > 0:
                return f"Page {best[1]}", _snip(best[2])
    except Exception:
        pass
    return "", ""
