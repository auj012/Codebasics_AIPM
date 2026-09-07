"""Turn each asset into searchable text.

One extractor per asset type - they all share the same extract() method, so the
indexer can treat them the same way (this is the polymorphism from the HLD).
"""
import re
from pathlib import Path


class TextExtractor:
    """Base class. Each subclass knows how to read ONE kind of asset."""

    def extract(self, path_or_url: str) -> str:
        raise NotImplementedError


class PdfExtractor(TextExtractor):
    def extract(self, path):
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            text = " ".join((pg.extract_text() or "") for pg in reader.pages)
            return " ".join(text.split())[:3000]
        except Exception:
            return ""


class PptxExtractor(TextExtractor):
    def extract(self, path):
        try:
            from pptx import Presentation
            prs = Presentation(str(path))
            chunks = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        chunks.append(shape.text_frame.text)
            return " ".join(" ".join(chunks).split())[:3000]
        except Exception:
            return ""


class ImageExtractor(TextExtractor):
    """v1: uses the FILENAME as text (thumbnails have descriptive names like
    'what_is_power_bi_thumbnail'). Reading text INSIDE the image (OCR) is a
    future upgrade - see README."""

    def extract(self, path):
        return Path(path).stem.replace("_", " ").replace("-", " ")


class YouTubeTranscriptExtractor(TextExtractor):
    """Fetch a YouTube video's transcript (with timestamps) ONCE, at index time.
    Falls back to an empty string if the video has no captions - the indexer then
    keeps just the title/description."""

    def fetch_segments(self, url):
        """Return [(start_seconds, text), ...] for a YouTube URL - [] if unavailable.
        The start time is what lets us jump to the exact moment later (Story B)."""
        try:
            from youtube_transcript_api import YouTubeTranscriptApi
            vid = self._video_id(url)
            if not vid:
                return []
            fetched = YouTubeTranscriptApi().fetch(vid, languages=["en"])
            return [(float(s.start), s.text) for s in fetched][:500]
        except Exception:
            return []

    def extract(self, url):
        return " ".join(t for _, t in self.fetch_segments(url))[:5000]

    @staticmethod
    def _video_id(url):
        m = re.search(r"(?:v=|youtu\.be/|/shorts/|/embed/)([A-Za-z0-9_-]{11})", str(url))
        return m.group(1) if m else ""
